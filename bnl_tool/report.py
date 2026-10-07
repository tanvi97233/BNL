"""SOP Section 7 newsletter and audit-log generation."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

from .fetch import Paper
from .screen import ScreeningDecision


def write_docx(path: Path, records: list[tuple[Paper, ScreeningDecision]], start_date, end_date, *, draft: bool) -> None:
    doc = Document()
    title = "DRAFT — Bi-Weekly Newsletter - Research and Review Papers" if draft else "Bi-Weekly Newsletter - Research and Review Papers"
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Date range covered: {start_date.isoformat()} to {end_date.isoformat()}")
    doc.add_paragraph(f"Total papers included: {len(records)}")
    doc.add_paragraph(f"Prepared on: {datetime.now().date().isoformat()}")
    for serial, (paper, decision) in enumerate(records, 1):
        doc.add_heading(f"{serial}. {paper.title}", level=1)
        values = {
            "Serial Number": str(serial),
            "Title": paper.title,
            "Disease Area": "; ".join(decision.disease_areas_matched),
            "Paper Type": decision.paper_type,
            "Summary": decision.summary,
            "Source Link": f"https://doi.org/{paper.doi}" if paper.doi else paper.url,
        }
        table = doc.add_table(rows=0, cols=2); table.style = "Table Grid"
        for label, value in values.items():
            row = table.add_row().cells; row[0].text = label
            if label == "Source Link": _add_hyperlink(row[1].paragraphs[0], value, value)
            else: row[1].text = value
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
    doc.save(path)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part; rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr"); color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); props.append(color); run.append(props)
    content = OxmlElement("w:t"); content.text = text; run.append(content); link.append(run); paragraph._p.append(link)


def write_audit_csv(path: Path, hits, papers: list[Paper], decisions: dict[str, ScreeningDecision]) -> None:
    fields = ["query_number", "query", "page", "rank", "title", "url", "snippet", "prefilter_reason", "google_results_url", "timestamp", "doi", "journal", "volume", "issue", "publication_date", "fetch_error", "include", "exclusion_reason", "paper_type", "domain", "disease_areas_matched", "summary", "final_relevance_confirmed"]
    by_url = {paper.url: paper for paper in papers}
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader()
        for hit in hits:
            paper = by_url.get(hit.url); d = decisions.get(hit.url)
            writer.writerow({"query_number": hit.query_number, "query": hit.query, "page": hit.page, "rank": hit.rank, "title": hit.title, "url": hit.url, "snippet": hit.snippet, "prefilter_reason": hit.prefilter_reason, "google_results_url": hit.google_results_url, "timestamp": hit.timestamp, "doi": paper.doi if paper else "", "journal": paper.journal if paper else "", "volume": paper.volume if paper else "", "issue": paper.issue if paper else "", "publication_date": paper.publication_date if paper else "", "fetch_error": paper.fetch_error if paper else "", "include": d.include if d else False, "exclusion_reason": d.exclusion_reason if d else hit.prefilter_reason, "paper_type": d.paper_type if d else "", "domain": d.domain if d else "", "disease_areas_matched": "; ".join(d.disease_areas_matched) if d else "", "summary": d.summary if d else "", "final_relevance_confirmed": d.final_relevance_confirmed if d else False})


def write_papers_exports(csv_path: Path, xlsx_path: Path, records: list[tuple[Paper, ScreeningDecision]]) -> None:
    rows = [{"Serial Number": index, "Title": paper.title, "Disease Area": "; ".join(d.disease_areas_matched), "Paper Type": d.paper_type, "Summary": d.summary, "Source Link": f"https://doi.org/{paper.doi}" if paper.doi else paper.url} for index, (paper, d) in enumerate(records, 1)]
    fields = ["Serial Number", "Title", "Disease Area", "Paper Type", "Summary", "Source Link"]
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
    wb = Workbook(); ws = wb.active; ws.title = "Papers"; ws.append(fields)
    for row in rows: ws.append([row[field] for field in fields])
    wb.save(xlsx_path)
